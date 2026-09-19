from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from .dataset import DatasetAdapter


class AttachmentReadError(RuntimeError):
    """Raised when an attachment cannot produce usable text."""


@dataclass(frozen=True)
class AttachmentContent:
    text: str
    spans: tuple[dict[str, Any], ...]


def read_attachment_text(adapter: DatasetAdapter, reference: str) -> str:
    """Extract text from a supported attachment referenced by an email."""
    return read_attachment_content(adapter, reference).text


def read_attachment_content(adapter: DatasetAdapter, reference: str) -> AttachmentContent:
    """Extract attachment text and best-effort source locations."""
    suffix = Path(reference).suffix.lower()
    try:
        if suffix == ".txt":
            content = _text_content(adapter.read_text(reference), "plain_text")
        elif suffix == ".pdf":
            content = _read_pdf(adapter, reference)
        elif suffix == ".docx":
            content = _text_content(_read_docx(adapter, reference), "docx_text")
        elif suffix == ".xlsx":
            content = _text_content(_read_xlsx(adapter, reference), "xlsx_text")
        else:
            raise AttachmentReadError(f"Unsupported attachment format: {suffix}")
    except AttachmentReadError:
        raise
    except Exception as error:
        raise AttachmentReadError(f"Could not read attachment: {reference}") from error

    text = content.text.strip()
    if not text:
        raise AttachmentReadError(f"Attachment contains no readable text: {reference}")
    return AttachmentContent(text=text, spans=_rebase_spans(content.spans, content.text))


def _read_pdf(adapter: DatasetAdapter, reference: str) -> AttachmentContent:
    from io import BytesIO
    from pypdf import PdfReader

    try:
        import fitz

        data = adapter.read_bytes(reference) if adapter.is_http else adapter.resolve_attachment(reference).read_bytes()
        document = fitz.open(stream=data, filetype="pdf")
        lines: list[str] = []
        spans: list[dict[str, Any]] = []
        offset = 0
        for page_number, page in enumerate(document, start=1):
            grouped: dict[tuple[int, int], list[tuple[float, float, float, float, str]]] = {}
            for word in page.get_text("words"):
                x0, y0, x1, y1, word_text, block, line, _ = word
                grouped.setdefault((block, line), []).append((x0, y0, x1, y1, word_text))
            for words in grouped.values():
                words.sort(key=lambda item: item[0])
                line_text = " ".join(item[4] for item in words).strip()
                if not line_text:
                    continue
                if lines:
                    offset += 1
                lines.append(line_text)
                spans.append({
                    "start": offset,
                    "end": offset + len(line_text),
                    "text": line_text,
                    "page": page_number,
                    "coordinates": {
                        "x0": min(item[0] for item in words),
                        "y0": min(item[1] for item in words),
                        "x1": max(item[2] for item in words),
                        "y1": max(item[3] for item in words),
                    },
                    "method": "pymupdf_text",
                })
                offset += len(line_text)
        text = "\n".join(lines).strip()
        document.close()
        if text:
            return AttachmentContent(text=text, spans=tuple(spans))
    except Exception:
        pass

    source = str(adapter.resolve_attachment(reference)) if not adapter.is_http else BytesIO(adapter.read_bytes(reference))
    try:
        reader = PdfReader(source)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if text.strip():
            return _text_content(text, "pdf_text")
    except Exception:
        pass
    return _text_content(_ocr_pdf_document(adapter, reference), "ocr")


def _ocr_pdf_document(adapter: DatasetAdapter, reference: str) -> str:
    """OCR a PDF using RapidOCR with pypdfium2 rasterization, falling back to pytesseract."""
    try:
        import pypdfium2 as pdfium
        from rapidocr_onnxruntime import RapidOCR

        engine = RapidOCR()
        data = adapter.read_bytes(reference) if adapter.is_http else adapter.resolve_attachment(reference).read_bytes()
        pdf = pdfium.PdfDocument(data)
        blocks: list[str] = []
        for page in pdf:
            arr = page.render(scale=200 / 72).to_numpy()
            result, _ = engine(arr)
            if result:
                page_text = "\n".join(line[1] for line in result if line[1].strip())
                if page_text:
                    blocks.append(page_text)
        text = "\n".join(blocks).strip()
        if text:
            return text
    except Exception as error:
        if "Data format error" in str(error) or "format error" in str(error).lower():
            raise AttachmentReadError(f"Corrupted PDF attachment: {reference}") from error

    # Secondary fallback: pytesseract if available
    try:
        from io import BytesIO
        from pypdf import PdfReader

        source = str(adapter.resolve_attachment(reference)) if not adapter.is_http else BytesIO(adapter.read_bytes(reference))
        reader = PdfReader(source)
        return _ocr_pdf_images(reader, reference)
    except AttachmentReadError:
        raise
    except Exception as error:
        raise AttachmentReadError(f"OCR failed for attachment: {reference}: {error}") from error


def _ocr_pdf_images(reader: object, reference: str) -> str:
    """OCR embedded PDF images when the PDF has no text layer."""
    from io import BytesIO

    try:
        import pytesseract
        from PIL import Image
    except ImportError as error:
        raise AttachmentReadError(
            "OCR requires the optional pytesseract and Pillow packages."
        ) from error

    pages = getattr(reader, "pages", [])
    blocks: list[str] = []
    try:
        for page in pages:
            for image in getattr(page, "images", []):
                with Image.open(BytesIO(image.data)) as source:
                    blocks.append(pytesseract.image_to_string(source))
    except Exception as error:
        raise AttachmentReadError(f"OCR failed for attachment: {reference}") from error
    text = "\n".join(blocks).strip()
    if not text:
        raise AttachmentReadError(f"OCR produced no text for attachment: {reference}")
    return text


def _read_docx(adapter: DatasetAdapter, reference: str) -> str:
    from io import BytesIO
    import docx

    source = BytesIO(adapter.read_bytes(reference)) if adapter.is_http else str(adapter.resolve_attachment(reference))
    document = docx.Document(source)
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                blocks.append(" | ".join(values))
    return "\n".join(blocks)


def _read_xlsx(adapter: DatasetAdapter, reference: str) -> str:
    from io import BytesIO
    from openpyxl import load_workbook

    source = BytesIO(adapter.read_bytes(reference)) if adapter.is_http else adapter.resolve_attachment(reference)
    workbook = load_workbook(
        filename=source,
        read_only=True,
        data_only=True,
    )
    blocks: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    blocks.append(" | ".join(values))
    finally:
        workbook.close()
    return "\n".join(blocks)


def _text_content(text: str, method: str) -> AttachmentContent:
    spans: list[dict[str, Any]] = []
    offset = 0
    for line in text.splitlines():
        if line.strip():
            spans.append({
                "start": offset,
                "end": offset + len(line),
                "text": line,
                "page": None,
                "coordinates": None,
                "method": method,
            })
        offset += len(line) + 1
    return AttachmentContent(text=text, spans=tuple(spans))


def _rebase_spans(spans: tuple[dict[str, Any], ...], original: str) -> tuple[dict[str, Any], ...]:
    shift = len(original) - len(original.lstrip())
    return tuple(
        {**span, "start": max(0, span["start"] - shift), "end": max(0, span["end"] - shift)}
        for span in spans
    )